# -*- coding: utf-8 -*-
"""把 GSE127465 独立复核的引用句插入两个投稿包正文 .docx。

插入位置：Discussion 的 "Biological interpretation" 段，紧接
"…and strong expression in CD8 T, CD4 T and γδ T cells." 之后、
"High expression does not imply strong genetic regulation:" 之前。

实现：定位包含锚点的 run，把新句插到该 run 内锚点之前——**不新建 run、不改变段落格式**，
其余文本一字不动。**先备份**。

用法：python 46_insert_replication_sentence_docx.py [--dry-run]
"""
import os
import shutil
import sys
from datetime import datetime

import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DRY = "--dry-run" in sys.argv

ANCHOR = "High expression does not imply strong genetic regulation:"
SENTENCE = ("The same compartmental pattern was reproduced in an independent treatment-na\u00efve "
            "NSCLC single-cell dataset (GSE127465; Zilionis et al., 54,773 cells; Additional files "
            "12\u201315), in which JAML remained myeloid-predominant (23\u201358% positive across "
            "myeloid and dendritic-cell subsets) and rare in epithelium (1.5%), CXADR was confined "
            "to the epithelial compartment (19.3% versus 0.5% in myeloid cells), and the "
            "epithelial-to-myeloid positivity ratio for JAML (0.05) was essentially identical to "
            "that in our discovery dataset (0.04); agreement between the expression matrix and the "
            "cell metadata was verified by per-cell mitochondrial-fraction concordance "
            "(Spearman \u03c1 = 1.0000 across 54,773 cells). ")


def process(pkg):
    path = os.path.join(ROOT, pkg, "01_Manuscript_main.docx")
    print(f"\n==== {pkg} ====")
    if not os.path.isfile(path):
        print("  [跳过] 找不到文件")
        return
    doc = docx.Document(path)
    target = None
    for p in doc.paragraphs:
        if ANCHOR in p.text:
            target = p
            break
    if target is None:
        print("  [跳过] 未找到锚点")
        return
    if "GSE127465" in target.text:
        print("  [跳过] 该段已包含 GSE127465（可能已插入过）")
        return

    hit = [r for r in target.runs if ANCHOR in r.text]
    if len(hit) != 1:
        print(f"  [跳过] 锚点落在 {len(hit)} 个 run 中，需人工处理（避免破坏格式）")
        return
    run = hit[0]
    pre = run.text.split(ANCHOR)[0]
    print(f"  锚点所在 run 之前的内容尾部：{pre[-60:]!r}")
    if DRY:
        print(f"  [dry-run] 将插入 {len(SENTENCE)} 字符")
        return

    bak = path + f".bak_{datetime.now():%Y%m%d_%H%M%S}"
    shutil.copy2(path, bak)
    print(f"  [备份] {os.path.relpath(bak, ROOT)}")
    run.text = pre + SENTENCE + ANCHOR + run.text.split(ANCHOR, 1)[1]
    doc.save(path)

    d2 = docx.Document(path)
    ok = any("GSE127465" in p.text for p in d2.paragraphs)
    print(f"  [校验] 段落中已含 GSE127465：{ok}")


def main():
    for pkg in ["BMCCancer投稿", "JTM投稿"]:
        process(pkg)
    print("\nDONE" + (" (dry-run)" if DRY else ""))


if __name__ == "__main__":
    main()
