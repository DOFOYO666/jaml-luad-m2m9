# -*- coding: utf-8 -*-
"""把 M2-M9 深化研究的中/英文 Markdown 稿转为投稿用 .docx。

格式遵循项目约定：
  - 正文 Times New Roman 11 pt（中文稿西文 Times New Roman、中文宋体）
  - 标题层级：Title / Heading 1 / Heading 2
  - 表格一律**三线表**（顶线、表头底线、底线；无竖线、无内部横线）
  - 图片单独放在 03_Figures/，正文只保留图题文字

用法：
  python 49_build_new_docx.py            # 生成中英两版
  python 49_build_new_docx.py --which en
"""
import io
import os
import re
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

PKG = r"C:\Users\86159\WorkBuddy\单细胞数据孟德尔随机化\M2-M9深化研究稿"

CFG = {
    "en": dict(src="manuscript_EN.md", out="01_Manuscript_EN.docx",
               latin="Times New Roman", ea="Times New Roman", size=11.0),
    "cn": dict(src="manuscript_CN.md", out="01_Manuscript_CN.docx",
               latin="Times New Roman", ea="宋体", size=11.0),
}


# ------------------------------------------------------------------ 样式工具
def set_font(run, latin, ea, size=None, bold=None, italic=None):
    run.font.name = latin
    run.font.size = Pt(size) if size else run.font.size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea)


def add_rich(par, text, latin, ea, size, base_bold=False):
    """解析 **bold** 与 *italic* 的行内标记，写入 paragraph。"""
    tokens = re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text)
    for tk in tokens:
        if not tk:
            continue
        if tk.startswith("**") and tk.endswith("**") and len(tk) > 4:
            r = par.add_run(tk[2:-2]); set_font(r, latin, ea, size, True, None)
        elif tk.startswith("*") and tk.endswith("*") and len(tk) > 2:
            r = par.add_run(tk[1:-1]); set_font(r, latin, ea, size, base_bold or None, True)
        else:
            r = par.add_run(tk); set_font(r, latin, ea, size, base_bold or None, None)


def clear_borders(table):
    tbl = table._element
    pr = tbl.tblPr
    for tag in ("w:tblBorders",):
        el = pr.find(qn(tag))
        if el is not None:
            pr.remove(el)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        b.append(e)
    pr.append(b)


def set_cell_border(cell, edge, sz=8):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    e = borders.find(qn(f"w:{edge}"))
    if e is None:
        e = OxmlElement(f"w:{edge}")
        borders.append(e)
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:color"), "000000")


def three_line(table):
    """三线表：顶线 + 表头底线 + 底线，无竖线/内部横线。"""
    clear_borders(table)
    n = len(table.rows)
    for c in table.rows[0].cells:
        set_cell_border(c, "top", 12)
        set_cell_border(c, "bottom", 8)
    for c in table.rows[n - 1].cells:
        set_cell_border(c, "bottom", 12)


# ------------------------------------------------------------------ 转换主体
def convert(md_path, out_path, cfg):
    latin, ea, size = cfg["latin"], cfg["ea"], cfg["size"]
    lines = io.open(md_path, encoding="utf-8").read().split("\n")
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = latin
    st.font.size = Pt(size)
    st.paragraph_format.space_after = Pt(6)
    # JTM/BMC: "Use double-line spacing" for the main manuscript text.
    st.paragraph_format.line_spacing = 2.0
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.2)

    # JTM/BMC: "Include line and page numbering".
    for s in d.sections:
        _sect_pr = s._sectPr
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        ln.set(qn("w:distance"), "360")
        cols = _sect_pr.find(qn("w:cols"))
        if cols is not None:                      # schema order: lnNumType precedes cols
            cols.addprevious(ln)
        else:
            _sect_pr.append(ln)

        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.line_spacing = 1.0
        run = fp.add_run()
        run.font.name = latin
        run.font.size = Pt(size - 1.0)
        for tag, attr in (("w:fldChar", ("w:fldCharType", "begin")),
                          ("w:instrText", ("xml:space", "preserve"))):
            el = OxmlElement(tag)
            el.set(qn(attr[0]), attr[1])
            if tag == "w:instrText":
                el.text = " PAGE "
            run._r.append(el)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(end)

    i = 0
    first_h1 = True
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()

        # 表格块
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [[c.strip() for c in r.strip("|").split("|")] for r in block]
            rows = [rows[0]] + rows[2:]                    # 去掉分隔行
            ncol = max(len(r) for r in rows)
            t = d.add_table(rows=0, cols=ncol)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, r in enumerate(rows):
                cells = t.add_row().cells
                for ci in range(ncol):
                    txt = r[ci] if ci < len(r) else ""
                    p = cells[ci].paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0
                    add_rich(p, txt, latin, ea, size - 1.0, base_bold=(ri == 0))
            three_line(t)
            d.add_paragraph()
            continue

        # 标题
        if s.startswith("# "):
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(0 if first_h1 else 18)
            p.paragraph_format.space_after = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if first_h1 else WD_ALIGN_PARAGRAPH.LEFT
            add_rich(p, s[2:], latin, ea, size + 5, base_bold=True)
            first_h1 = False
            i += 1
            continue
        if s.startswith("## "):
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            add_rich(p, s[3:], latin, ea, size + 2.5, base_bold=True)
            i += 1
            continue
        if s.startswith("### "):
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            add_rich(p, s[4:], latin, ea, size + 1, base_bold=True)
            i += 1
            continue

        if s in ("---", "***", "___"):
            p = d.add_paragraph()
            r = p.add_run(); set_font(r, latin, ea, size)
            r.add_break(WD_BREAK.LINE)
            i += 1
            continue

        if not s:
            i += 1
            continue

        # 有序/无序列表
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            p = d.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, m.group(2), latin, ea, size)
            i += 1
            continue
        if re.match(r"^[-*]\s+\S", s) and not s.startswith("* " * 2):
            p = d.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, s[2:], latin, ea, size)
            i += 1
            continue

        p = d.add_paragraph()
        add_rich(p, s, latin, ea, size)
        i += 1

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    d.save(out_path)
    return out_path


def verify(path, md_path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    txt = "\n".join(parts)
    md = io.open(md_path, encoding="utf-8").read()
    keys = [
        ("含 0.757（M2 髓系 DC）", "0.757" in txt),
        ("含 0.05（复发集比值）", "0.05" in txt),
        ("含 6.55（配对比值）", "6.55" in txt),
        ("含 0.881（ICB）", "0.881" in txt),
        ("含 0.1275 或 0.128（拟时序）", ("0.1275" in txt) or ("0.128" in txt)),
        ("含 1.0000（对齐 ρ）", "1.0000" in txt),
        ("不含 'up to ~3.0'", "up to ~3.0" not in txt),
    ]
    print(f"  {os.path.basename(path)}：段落 {len(d.paragraphs)}，表格 {len(d.tables)}")
    for k, v in keys:
        print(f"    {'OK  ' if v else '!!  '}{k}: {v}")
    n_tbl_md = len(re.findall(r"(?m)^\|", md))
    print(f"    markdown 表格行数 {n_tbl_md}；docx 表格数 {len(d.tables)}")


if __name__ == "__main__":
    which = sys.argv[sys.argv.index("--which") + 1] if "--which" in sys.argv else None
    for k, cfg in CFG.items():
        if which and k != which:
            continue
        src = os.path.join(PKG, cfg["src"])
        dst = os.path.join(PKG, cfg["out"])
        convert(src, dst, cfg)
        print(f"=== {k.upper()} ===")
        verify(dst, src)
        print()
    print("DONE")
