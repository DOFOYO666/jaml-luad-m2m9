# -*- coding: utf-8 -*-
"""Render JTM投稿_M2M9/manuscript_JTM.md to 01_Manuscript_JTM.docx.

Reuses the three-line-table / inline-markup machinery from scripts/49_build_new_docx.py
(imported by path, so the two builders can never drift apart), then verifies the output
by reading the paragraphs *and* the table cells back.

Run: python scripts/65_build_jtm_docx.py
"""
import importlib.util
import io
import os
import re
import sys

import docx

sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "JTM投稿_M2M9")
MD = os.path.join(OUT, "manuscript_JTM.md")
DOCX = os.path.join(OUT, "01_Manuscript_JTM.docx")
BUILDER = os.path.join(BASE, "scripts", "49_build_new_docx.py")

spec = importlib.util.spec_from_file_location("builder49", BUILDER)
b49 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(b49)

CFG = dict(latin="Times New Roman", ea="Times New Roman", size=11.0)


def verify():
    d = docx.Document(DOCX)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    txt = "\n".join(parts)
    md = io.open(MD, encoding="utf-8").read()

    keys = [
        ("title page: corresponding author", "xingxinghuoshu@163.com" in txt),
        ("title page: ORCID", "0009-0007-8106-0425" in txt),
        ("abstract headings present",
         all(k in txt for k in ("Background.", "Methods.", "Results.", "Conclusions."))),
        ("keyword count <= 10",
         len([k for k in md.split("**Keywords:**")[1].split("\n")[0].split(";")]) <= 10),
        ("key number 0.757 (myeloid DC)", "0.757" in txt),
        ("key number 6.55 (paired ratio)", "6.55" in txt),
        ("key number 0.881 (ICB null)", "0.881" in txt),
        ("alignment rho 1.0000", "1.0000" in txt),
        ("additional files named per BMC rule", "Additional file 17" in txt),
        ("no stale 'Supplementary Material' citation", "Supplementary Material" not in txt),
        ("Abbreviations section present", "Abbreviations" in txt),
        ("Acknowledgements sub-heading present", "Acknowledgements." in txt),
        ("mandated Availability opening", "The dataset(s) supporting the conclusions of this article are available in" in txt),
        ("software fields present", "Archived version:" in txt and "Operating system(s):" in txt
         and "Programming language:" in txt and "Any restrictions to use by non-academics:" in txt),
        ("no stale 'target triage' keyword", "target triage" not in txt),
        ("section: Additional files", "Additional files" in txt),
        ("section: Declarations", "Declarations" in txt),
        ("section: References before Figure legends",
         txt.find("References") < txt.find("Figure legends")),
        ("section: Figure legends before Tables",
         txt.find("Figure legends") < txt.find("Table 1.")),
        ("5 tables", len(d.tables) == 5),
    ]
    print("  %s: paragraphs %d, tables %d" % (os.path.basename(DOCX), len(d.paragraphs), len(d.tables)))
    bad = 0
    for k, v in keys:
        print("    %s%s: %s" % ("OK  " if v else "!!  ", k, v))
        bad += 0 if v else 1
    n_md_rows = len(re.findall(r"(?m)^\|", md))
    print("    markdown table rows %d ; docx tables %d" % (n_md_rows, len(d.tables)))
    return bad


def main():
    b49.convert(MD, DOCX, CFG)
    bad = verify()
    print("\n%s" % ("DONE" if bad == 0 else "DONE WITH %d FAILED CHECK(S)" % bad))


if __name__ == "__main__":
    main()
